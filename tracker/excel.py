"""Write the Excel workbook."""
import os
from datetime import date, timedelta

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

import config

HEADERS = ["First seen", "Type", "Title", "Organisation", "Location", "Deadline / Start",
           "Days left", "Pay / Prize", "Source", "Link", "Match", "Status", "Notes", "ID"]
WIDTHS = [11, 11, 55, 26, 30, 14, 9, 24, 18, 12, 7, 11, 45, 1]
HEAD_FILL = PatternFill("solid", fgColor="1F4E78")
NEW_FILL = PatternFill("solid", fgColor="E2EFDA")
STATUSES = "New,Interested,Applied,Interview,Rejected,Skipped"


def _old_statuses(path):
    """Keep the Status you typed in the previous file."""
    if not os.path.exists(path):
        return {}
    try:
        ws = load_workbook(path, read_only=True)["All"]
        return {r[13]: r[11] for r in ws.iter_rows(min_row=2, values_only=True) if r and r[13] and r[11]}
    except Exception:
        return {}


def _days_left(d):
    try:
        return (date.fromisoformat(d) - date.today()).days
    except (ValueError, TypeError):
        return None


def _sheet(wb, name, rows, statuses, today):
    ws = wb.create_sheet(name)
    ws.append(HEADERS)
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = HEAD_FILL
    for r in rows:
        ws.append([r["first_seen"], r["type"], r["title"], r["org"], r["location"], r["deadline"],
                   _days_left(r["deadline"]), r["pay"], r["source"], "Open",
                   int(r["score"] or 0), statuses.get(r["id"], "New"), r["notes"], r["id"]])
        row = ws.max_row
        ws.cell(row, 10).hyperlink = r["link"]
        ws.cell(row, 10).font = Font(color="0563C1", underline="single")
        if r["first_seen"] == today:
            for c in ws[row][:3]:
                c.fill = NEW_FILL
    for i, w in enumerate(WIDTHS):
        ws.column_dimensions[chr(65 + i)].width = w
    ws.column_dimensions["N"].hidden = True
    ws.freeze_panes = "D2"
    ws.auto_filter.ref = ws.dimensions
    dv = DataValidation(type="list", formula1=f'"{STATUSES}"', allow_blank=True)
    ws.add_data_validation(dv)
    dv.add(f"L2:L{max(ws.max_row, 2)}")


def write(path, rows, health):
    statuses = _old_statuses(path)
    today = date.today().isoformat()
    soon = (date.today() + timedelta(days=config.CLOSING_SOON_DAYS)).isoformat()
    wb = Workbook()
    wb.remove(wb.active)
    tabs = [
        ("NEW TODAY", [r for r in rows if r["first_seen"] == today]),
        ("Internships", [r for r in rows if r["type"] == "Internship"]),
        ("Jobs", [r for r in rows if r["type"] == "Job"]),
        ("Hackathons", [r for r in rows if r["type"] == "Hackathon"]),
        ("Contests", [r for r in rows if r["type"] == "Contest"]),
        ("Closing Soon", sorted([r for r in rows if r["deadline"] and today <= r["deadline"] <= soon],
                                key=lambda r: r["deadline"])),
        ("All", rows),
    ]
    for name, sub in tabs:
        _sheet(wb, name, sub, statuses, today)

    ws = wb.create_sheet("Source Health")
    ws.append(["Source", "Status", "Rows", "Detail"])
    for c in ws[1]:
        c.font = Font(bold=True)
    for h in health:
        ws.append(list(h))
    ws.column_dimensions["A"].width = 30
    ws.column_dimensions["D"].width = 120

    tmp = path + ".tmp.xlsx"
    wb.save(tmp)
    os.replace(tmp, path)  # fails if the file is open in Excel; caller reports it
